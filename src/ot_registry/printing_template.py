"""Печать по DOCX-шаблонам пользователя: сборка документа."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document

from . import config, docx_templates
from .errors import ValidationError
from .excel_store import Book
from .printing import (
    PrintResult,
    _convert_pdf,
    _output_dir,
    placeholder_values,
)


def _count_cards(document) -> int:
    """Сколько карточек в шаблоне: по числу вхождений плейсхолдера fio."""
    total = 0
    for text in docx_templates._iter_text(document):
        total += len(re.findall(r"\{\{\s*fio\s*\}\}", text))
    return max(total, 1)


def _fill_one(document, values: dict[str, str]) -> None:
    docx_templates._replace_in_document(document, values)


def _fill_many(document, values: list[dict[str, str]]) -> None:
    """Шаблон с несколькими карточками: подставляем значения по порядку."""
    state = {"index": -1, "card": values[0] if values else {}}

    def repl(match: re.Match) -> str:
        name = match.group(1)
        if name == "fio":
            state["index"] += 1
            idx = min(state["index"], len(values) - 1)
            state["card"] = values[idx]
        card = state["card"]
        return card.get(name, match.group(0))

    pattern = r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}"
    for paragraph in docx_templates._iter_text_paragraphs(document):
        if "{{" not in paragraph.text:
            continue
        new_text = re.sub(pattern, repl, paragraph.text)
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""


def build_document(book: Book, rows: list[dict[str, Any]], *,
                   template_path: Path, manual: bool) -> PrintResult:
    """Заполнить шаблон и сохранить DOCX (+ PDF, если есть конвертер)."""
    document = Document(template_path)
    per_template = _count_cards(document)
    values = [
        placeholder_values(book, row, card_number=index + 1)
        for index, row in enumerate(rows)
    ]

    if per_template == 1:
        # Обычный случай: одна карточка в шаблоне. Заполняем первый лист,
        # остальные записи добавляем на новые страницы с тем же шаблоном.
        _fill_one(document, values[0])
        for offset in range(per_template, len(values)):
            document.add_page_break()
            extra = Document(template_path)
            _fill_one(extra, values[offset])
            _append_document(document, extra)
    else:
        _fill_many(document, values[:per_template])
        rest = values[per_template:]
        while rest:
            chunk, rest = rest[:per_template], rest[per_template:]
            document.add_page_break()
            extra = Document(template_path)
            if len(chunk) == 1:
                _fill_one(extra, chunk[0])
            else:
                _fill_many(extra, chunk)
            _append_document(document, extra)

    out_dir = _output_dir(book)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    prefix = "manual" if manual else "filled"
    docx_path = out_dir / f"{prefix}_tpl_{stamp}_{len(rows)}.docx"
    document.save(docx_path)

    pdf_path = _convert_pdf(docx_path)
    note = ""
    if not pdf_path:
        note = ("PDF не создан: не найден LibreOffice или Word. "
                "Файл DOCX готов к печати.")
    return PrintResult(docx_path=docx_path, pdf_path=pdf_path,
                       cards=len(rows), note=note)


def _append_document(target, source) -> None:
    """Добавить содержимое одного документа в конец другого.

    Копируются абзацы и таблицы с сохранением текста и базового оформления.
    Сложная вёрстка (колонтитулы, разрывы разделов) не переносится —
    для бланков это не требуется.
    """
    from copy import deepcopy

    body = target.element.body
    for child in source.element.body.iterchildren():
        # Пропускаем служебный элемент sectPr: он описывает параметры
        # последнего раздела исходного документа.
        if child.tag.endswith("sectPr"):
            continue
        body.append(deepcopy(child))


def _extract_rows(book: Book, training_ids: Iterable[str]) -> list[dict[str, Any]]:
    ids = [i for i in training_ids if i]
    if not ids:
        raise ValidationError("Выберите хотя бы одну запись для печати.")
    rows = [r for r in book.table("Инструктажи").rows if str(r.get("id") or "") in set(ids)]
    if not rows:
        raise ValidationError("Выбранные записи не найдены в книге.")
    # Сохраняем порядок, в котором записи выбирались.
    order = {value: index for index, value in enumerate(ids)}
    return sorted(rows, key=lambda r: order.get(str(r.get("id") or ""), 10**6))


def print_filled_template(book: Book, training_ids: Iterable[str], *,
                          cards_per_page: int = 1) -> PrintResult | None:
    """Печать заполненных бланков по шаблону. None — если шаблона нет."""
    template_path = docx_templates.pick_template(cards_per_page)
    if template_path is None:
        return None
    rows = _extract_rows(book, training_ids)
    result = build_document(book, rows, template_path=template_path, manual=False)

    from .journal import log_event

    log_event(book.logs_dir(), organization=book.name, action="print_filled",
              entity="Инструктажи",
              details={"ids": [str(r.get("id")) for r in rows],
                       "template": template_path.name, "cards": len(rows)})
    return result


def print_blank_template(book: Book, *, training_type: str, count: int = 1,
                         cards_per_page: int = 1) -> PrintResult | None:
    """Печать пустых бланков по шаблону: данные не подставляются."""
    template_path = docx_templates.pick_template(cards_per_page)
    if template_path is None:
        return None
    if training_type not in config.TRAINING_TYPES:
        raise ValidationError("Выберите вид инструктажа для пустых бланков.")
    count = max(1, min(int(count or 1), 50))

    empty = {
        "code": config.MANUAL_CODE,
        "type": training_type,
        "fio": "",
        "position": "",
        "department": "",
        "birth_date": "",
        "program": "",
        "instructor_fio": "",
        "plan_date": "",
        "fact_date": "",
    }
    rows = [dict(empty) for _ in range(count)]
    result = build_document(book, rows, template_path=template_path, manual=True)

    from .journal import log_event

    log_event(book.logs_dir(), organization=book.name, action="print_blank",
              entity="Печать",
              details={"type": training_type, "count": count,
                       "template": template_path.name})
    return result
