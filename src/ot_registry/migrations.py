"""Миграции формата книги Excel.

Схема книги версионируется. Перед изменением существующей книги создаётся
резервная копия; понижение версии не выполняется. Неизвестная более новая
версия открывается только для чтения.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import config
from .errors import ConfigError
from .excel_store import BODY_FONT, BORDER, DATE_FORMAT, DATE_FORMAT as _DF, fingerprint
from .schema import SHEET_BY_NAME, SHEETS


@dataclass
class MigrationResult:
    from_version: int
    to_version: int
    actions: list[str]
    backup: str = ""


def book_schema_version(path: Path) -> int:
    """Прочитать версию схемы из листа «Организация» без полной загрузки."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Организация" not in wb.sheetnames:
            return 0
        ws = wb["Организация"]
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if row and str(row[0]).strip() == "Версия схемы":
                try:
                    return int(str(row[1]).strip())
                except (TypeError, ValueError):
                    return 0
        return 0
    finally:
        wb.close()


def _backup(path: Path) -> Path:
    folder = path.parent / "backups"
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"before_migration_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    import shutil

    shutil.copy2(path, target)
    return target


def _set_org_value(wb, key: str, value) -> None:
    ws = wb["Организация"]
    for row_idx in range(2, ws.max_row + 1):
        if str(ws.cell(row=row_idx, column=1).value or "").strip() == key:
            ws.cell(row=row_idx, column=2, value=value)
            return
    for row_idx in range(2, ws.max_row + 2):
        if not ws.cell(row=row_idx, column=1).value:
            ws.cell(row=row_idx, column=1, value=key)
            ws.cell(row=row_idx, column=2, value=value)
            return


def _ensure_columns(path: Path, actions: list[str]) -> None:
    """Добавить недостающие колонки справа, не сдвигая существующие данные."""
    wb = load_workbook(path)
    try:
        for sheet in SHEETS:
            if sheet.name not in wb.sheetnames:
                actions.append(f"лист «{sheet.name}» отсутствует — требуется восстановление")
                continue
            ws = wb[sheet.name]
            existing = {
                str(ws.cell(row=1, column=idx).value or "").strip()
                for idx in range(1, ws.max_column + 1)
            }
            next_col = ws.max_column + 1
            for column in sheet.columns:
                if column.title in existing:
                    continue
                cell = ws.cell(row=1, column=next_col, value=column.title)
                header_src = ws.cell(row=1, column=1)
                if header_src.has_style:
                    cell._style = copy(header_src._style)
                ws.column_dimensions[get_column_letter(next_col)].width = column.width
                if column.type == "date":
                    for row_idx in range(2, ws.max_row + 1):
                        ws.cell(row=row_idx, column=next_col).number_format = DATE_FORMAT
                actions.append(f"«{sheet.name}»: добавлена колонка «{column.title}»")
                next_col += 1
        _set_org_value(wb, "Версия схемы", config.SCHEMA_VERSION)
        wb.save(path)
    finally:
        wb.close()


def migrate(path: Path) -> MigrationResult:
    """Привести книгу к текущей версии схемы.

    Безопасно для пустых и уже актуальных книг. Для новых обязательных
    реквизитов данные не выдумываются: добавляются только колонки, а
    заполнение контролируется проверками при регистрации.
    """
    path = Path(path)
    if not path.exists():
        raise ConfigError(f"Файл книги не найден: {path}")

    current = book_schema_version(path)
    if current > config.SCHEMA_VERSION:
        raise ConfigError(
            f"Книга создана в более новой версии программы (схема {current}). "
            f"Текущая версия поддерживает схему {config.SCHEMA_VERSION}. "
            "Обновите приложение — запись в эту книгу заблокирована."
        )

    if current == config.SCHEMA_VERSION:
        return MigrationResult(current, current, [])

    backup = _backup(path)
    actions: list[str] = []
    _ensure_columns(path, actions)

    from .journal import log_event

    log_event(path.parent / "logs", organization=path.stem, action="migration",
              entity="Excel", reason=f"схема {current} → {config.SCHEMA_VERSION}",
              details={"actions": actions, "backup": backup.name})
    return MigrationResult(current, config.SCHEMA_VERSION, actions, str(backup))
