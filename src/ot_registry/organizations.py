"""Список подключённых организаций (книг Excel).

Хранится в `settings/organizations.json`. Хранит только путь и отображаемые
поля; данные людей и инструктажей остаются в книгах.
"""

from __future__ import annotations

import json
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from . import config
from .errors import ConfigError
from .schema import ORG_KEY

CATALOG_FILE = "organizations.json"
CATALOG_VERSION = 1


@dataclass
class OrganizationEntry:
    uuid: str
    name: str
    path: str
    org_uuid: str = ""
    created_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))
    last_opened_at: str = ""
    color: str = "#1F4E79"

    def resolve(self, base: Path | None = None) -> Path:
        p = Path(self.path)
        if not p.is_absolute():
            p = (base or config.data_dir()) / p
        return p


class Catalog:
    """Набор организаций и операции над ним."""

    def __init__(self, entries: list[OrganizationEntry] | None = None,
                 *, base_dir: Path | None = None) -> None:
        self.entries: list[OrganizationEntry] = entries or []
        self.base_dir = base_dir or config.data_dir()

    # -- файл ------------------------------------------------------------
    @property
    def path(self) -> Path:
        return self.base_dir / "settings" / CATALOG_FILE

    @classmethod
    def load(cls, base_dir: Path | None = None) -> "Catalog":
        base = base_dir or config.data_dir()
        path = base / "settings" / CATALOG_FILE
        if not path.exists():
            return cls([], base_dir=base)
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ConfigError(
                f"Не удалось прочитать список организаций ({path}). "
                "Проверьте файл или восстановите его из резервной копии."
            ) from exc
        entries = []
        for item in raw.get("organizations", []):
            try:
                entries.append(OrganizationEntry(**item))
            except TypeError:
                continue
        return cls(entries, base_dir=base)

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        payload = {
            "version": CATALOG_VERSION,
            "organizations": [asdict(e) for e in self.entries],
        }
        tmp = self.path.with_suffix(".tmp")
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        tmp.replace(self.path)

    # -- выборка ---------------------------------------------------------
    def get(self, org_uuid: str) -> OrganizationEntry | None:
        for entry in self.entries:
            if entry.uuid == org_uuid:
                return entry
        return None

    def require(self, org_uuid: str) -> OrganizationEntry:
        entry = self.get(org_uuid)
        if entry is None:
            raise ConfigError(
                "Организация не найдена. Вернитесь на главную страницу и выберите организацию."
            )
        return entry

    def by_book_uuid(self, book_uuid: str) -> OrganizationEntry | None:
        for entry in self.entries:
            if entry.org_uuid and entry.org_uuid == book_uuid:
                return entry
        return None

    def existing_paths(self) -> set[str]:
        return {str(e.resolve(self.base_dir)) for e in self.entries}

    # -- изменение -------------------------------------------------------
    def add(self, name: str, path: Path, *, org_uuid: str | None = None,
            color: str | None = None) -> OrganizationEntry:
        entry = OrganizationEntry(
            uuid=uuid.uuid4().hex,
            name=name,
            path=str(path),
            org_uuid=org_uuid or "",
            color=color or self._next_color(),
        )
        self.entries.append(entry)
        self.save()
        return entry

    def remove(self, org_uuid: str, *, delete_files: bool = False) -> None:
        entry = self.get(org_uuid)
        if entry is None:
            return
        self.entries = [e for e in self.entries if e.uuid != org_uuid]
        self.save()
        if delete_files:
            folder = entry.resolve(self.base_dir).parent
            # Удаляем только каталог книги внутри organizations/, не произвольный путь.
            try:
                if config.organizations_dir() in folder.parents:
                    import shutil

                    shutil.rmtree(folder)
            except OSError:
                pass

    def touch(self, org_uuid: str) -> None:
        entry = self.get(org_uuid)
        if entry:
            entry.last_opened_at = datetime.now().isoformat(timespec="seconds")
            self.save()

    def _next_color(self) -> str:
        palette = ["#1F4E79", "#2E7D32", "#6A1B9A", "#B26A00", "#00695C", "#8E244D"]
        return palette[len(self.entries) % len(palette)]

    def __len__(self) -> int:
        return len(self.entries)

    def __iter__(self):
        return iter(self.entries)


def prepare_org_dir(*, base_dir: Path | None = None) -> tuple[Path, Path]:
    """Создать каталог организации и вернуть (папка, путь книги).

    Имя каталога — случайный идентификатор, а не название организации:
    переименование или недопустимые символы в названии не влияют на путь.
    """
    base = base_dir or config.organizations_dir()
    slug = uuid.uuid4().hex
    folder = base / slug
    folder.mkdir(parents=True, exist_ok=True)
    return folder, folder / "registry.xlsx"


def read_org_meta(path: Path) -> dict[str, str]:
    """Прочитать только служебные значения книги (без полной загрузки таблиц)."""
    from openpyxl import load_workbook

    if not Path(path).exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Организация" not in wb.sheetnames:
            return {}
        ws = wb["Организация"]
        meta: dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                meta[str(row[0]).strip()] = "" if row[1] is None else str(row[1])
        return meta
    finally:
        wb.close()


def org_display_name(meta: dict[str, Any]) -> str:
    return str(meta.get(ORG_KEY["name"], "") or "")
