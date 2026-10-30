"""Генерация книги Excel по шаблону без запуска приложения.

Назначение:
- проверить, как выглядит книга с формулами и защитой;
- создать заготовку для новой организации;
- подготовить демонстрационный файл для тестирования и обучения.

Примеры:
    python scripts/make_excel.py --demo --out demo_Слайс_ЛТД.xlsx
    python scripts/make_excel.py --name "ООО Ромашка" --inn 7700111222 --out romashka.xlsx
    python scripts/make_excel.py --demo --no-protect --out открытая.xlsx
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ot_registry import config, excel_store, template  # noqa: E402
from ot_registry.schema import ORG_KEY, SHEETS  # noqa: E402


def describe(path: Path) -> None:
    """Краткий отчёт о созданной книге: что внутри."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    print(f"\nФайл: {path}")
    print(f"Размер: {path.stat().st_size / 1024:.0f} КБ")
    print(f"Листов: {len(wb.sheetnames)}")

    for sheet in SHEETS:
        if sheet.name not in wb.sheetnames:
            continue
        ws = wb[sheet.name]
        rows = max(ws.max_row - 1, 0) if sheet.kind == "data" else 0
        formulas = sum(
            1 for row in ws.iter_rows() for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        )
        parts = [f"строк: {rows}"]
        if formulas:
            parts.append(f"формул: {formulas}")
        if ws.protection.sheet:
            parts.append("защищён")
        print(f"  {sheet.name:<22} {', '.join(parts)}")

    password = excel_store._read_protection_password(wb)
    if password:
        print(f"\nПароль защиты структуры: {password}")
        print("Нужен, чтобы снять защиту листа в Excel вручную.")
    else:
        print("\nЗащита структуры не включена.")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Создать книгу Excel по шаблону instrOT",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("--name", default="Демонстрационная организация",
                        help="краткое название организации")
    parser.add_argument("--full-name", default="", help="полное наименование")
    parser.add_argument("--inn", default="", help="ИНН")
    parser.add_argument("--kpp", default="", help="КПП")
    parser.add_argument("--ogrn", default="", help="ОГРН")
    parser.add_argument("--address", default="", help="адрес")
    parser.add_argument("--director", default="", help="руководитель")
    parser.add_argument("--ot-specialist", default="", help="специалист по охране труда")
    parser.add_argument("--demo", action="store_true",
                        help="добавить демонстрационные данные")
    parser.add_argument("--no-protect", action="store_true",
                        help="не защищать структуру паролем")
    parser.add_argument("--out", type=Path, default=None,
                        help="путь к итоговому файлу (по умолчанию в текущем каталоге)")
    args = parser.parse_args(argv)

    name = args.name
    if args.demo and args.name == "Демонстрационная организация":
        name = "ООО «Слайс ЛТД»"

    out = args.out or Path(f"{name.replace(' ', '_').replace(chr(171), '').replace(chr(187), '')}.xlsx")
    if out.exists():
        print(f"Файл уже существует: {out}\nУкажите другой путь через --out.", file=sys.stderr)
        return 1

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / "registry.xlsx"
        book = excel_store.create_book(
            tmp_path,
            data_dir=Path(tmp),
            name=name,
            full_name=args.full_name or name,
            inn=args.inn,
            kpp=args.kpp,
            ogrn=args.ogrn,
            address=args.address,
            director=args.director,
            ot_specialist=args.ot_specialist,
            with_demo_data=args.demo,
        )
        if args.no_protect:
            book.protection_password = ""
            excel_store.save_book(
                book,
                action="regenerate_unprotected",
                reason="Книга создана без защиты структуры",
                make_backup=False,
                protect=False,
            )
        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(book.path, out)

    describe(out)
    print("\nГотово. Файл можно открыть в Excel или подключить в приложении.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
