"""Автотесты шаблона книги, формул, защиты и демонстрационных данных."""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from ot_registry import config, demo_data, excel_store, manual_mode, template  # noqa: E402
from ot_registry.schema import SHEETS, column_letter  # noqa: E402


@pytest.fixture()
def plain_book(tmp_path):
    return excel_store.create_book(
        tmp_path / "org" / "registry.xlsx", data_dir=tmp_path / "org", name="Тест"
    )


@pytest.fixture()
def demo_book(tmp_path):
    return excel_store.create_book(
        tmp_path / "demo" / "registry.xlsx",
        data_dir=tmp_path / "demo",
        name="ООО Тест",
        full_name="ООО «Тест»",
        with_demo_data=True,
    )


class TestTemplateStructure:
    def test_all_sheets_created(self, plain_book):
        wb = load_workbook(plain_book.path)
        assert wb.sheetnames == [s.name for s in SHEETS]

    def test_headers_match_schema(self, plain_book):
        wb = load_workbook(plain_book.path)
        for sheet in SHEETS:
            if sheet.kind != "data":
                continue
            ws = wb[sheet.name]
            titles = [ws.cell(row=1, column=idx).value for idx in range(1, len(sheet.columns) + 1)]
            assert titles == [col.title for col in sheet.columns], sheet.name

    def test_formula_headers_present(self, plain_book):
        wb = load_workbook(plain_book.path)
        ws = wb["Инструктажи"]
        count = len(next(s for s in SHEETS if s.name == "Инструктажи").columns)
        headers = [ws.cell(row=1, column=idx).value for idx in range(count + 1, count + 3)]
        assert headers == ["Дней до срока", "Состояние"]

    def test_formulas_written_for_rows(self, demo_book):
        wb = load_workbook(demo_book.path)
        ws = wb["Инструктажи"]
        count = len(next(s for s in SHEETS if s.name == "Инструктажи").columns)
        rows = ws.max_row
        assert rows >= 3
        for row in (2, 3, rows):
            cell = ws.cell(row=row, column=count + 1)
            assert isinstance(cell.value, str) and cell.value.startswith("=IF(")

    def test_no_semicolon_separators(self, demo_book):
        """Формулы должны использовать запятые: Excel и LibreOffice читают их одинаково."""
        wb = load_workbook(demo_book.path)
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        assert ";" not in cell.value, f"{name}!{cell.coordinate}"

    def test_summary_sheet_has_formulas(self, plain_book):
        wb = load_workbook(plain_book.path)
        ws = wb["Сводка"]
        formulas = [
            ws.cell(row=row, column=2).value
            for row in range(2, 20)
            if isinstance(ws.cell(row=row, column=2).value, str)
            and ws.cell(row=row, column=2).value.startswith("=")
        ]
        assert len(formulas) >= 10

    def test_instruction_sheet_has_text(self, plain_book):
        wb = load_workbook(plain_book.path)
        ws = wb["Инструкция"]
        text = " ".join(str(ws.cell(row=r, column=1).value or "") for r in range(1, 40))
        assert "КАК РАБОТАТЬ" in text
        assert "парол" in text.lower()

    def test_data_validation_on_status(self, plain_book):
        wb = load_workbook(plain_book.path)
        validations = wb["Инструктажи"].data_validations.dataValidation
        assert validations, "нет проверки данных на статусе"
        joined = " ".join(str(v.formula1) for v in validations)
        assert "Запланирован" in joined


class TestProtection:
    def test_book_is_protected(self, plain_book):
        wb = load_workbook(plain_book.path)
        assert all(ws.protection.sheet for ws in wb.worksheets)
        assert wb.security.lockStructure

    def test_insert_rows_blocked(self, plain_book):
        wb = load_workbook(plain_book.path)
        prot = wb["Работники"].protection
        assert prot.insertRows and prot.deleteRows and prot.insertColumns

    def test_password_stored_and_readable(self, plain_book):
        wb = load_workbook(plain_book.path)
        password = excel_store._read_protection_password(wb)
        assert password
        reopened = excel_store.open_book(plain_book.path, data_dir=plain_book.data_dir)
        assert reopened.protection_password == password

    def test_saving_keeps_protection(self, plain_book):
        plain_book.tables["Работники"].add(
            {"id": "x", "fio": "Тестов Тест", "position": "Дворник", "status": "Работает"}
        )
        excel_store.save_book(plain_book, action="test")
        wb = load_workbook(plain_book.path)
        assert wb["Работники"].protection.sheet

    def test_manual_mode_removes_and_restores_protection(self, plain_book):
        assert load_workbook(plain_book.path)["Инструктажи"].protection.sheet
        manual_mode.open_for_manual(plain_book)
        assert not load_workbook(plain_book.path)["Инструктажи"].protection.sheet

        wb = load_workbook(plain_book.path)
        wb["Работники"].append(
            [None, "1", "Новый Работник", "01.01.1990", "Дворник", "", "", "", "", "Работает", "6", "", ""]
        )
        wb.save(plain_book.path)

        manual_mode.check_book(plain_book)
        manual_mode.accept_changes(plain_book, reason="тест")
        assert load_workbook(plain_book.path)["Инструктажи"].protection.sheet


class TestDictionaries:
    def test_default_dictionaries_filled(self, plain_book):
        kinds = {row["kind"] for row in plain_book.table("Справочники").rows}
        assert {"position", "department", "program", "reason"} <= kinds

    def test_dictionary_values_reach_interface(self, plain_book):
        from ot_registry import services

        dicts = services.dictionaries(plain_book)
        assert "Дворник" in dicts["position"]
        assert "Участок озеленения" in dicts["department"]


class TestDemoData:
    def test_demo_rows_created(self, demo_book):
        assert len(demo_book.table("Работники").rows) == len(demo_data.DEMO_PERSONS)
        assert len(demo_book.table("Инструктирующие").rows) == len(demo_data.DEMO_INSTRUCTORS)
        assert len(demo_book.table("Инструктажи").rows) > 20

    def test_every_worker_has_birth_date(self, demo_book):
        for row in demo_book.table("Работники").rows:
            assert isinstance(row["birth_date"], date)

    def test_completed_trainings_have_numbers(self, demo_book):
        completed = [r for r in demo_book.table("Инструктажи").rows if r["status"] == "Проведён"]
        assert completed
        numbers = set()
        for row in completed:
            assert row["number_ot"].startswith("ОТ")
            assert row["number_type"]
            assert row["case_id"] and row["pack_id"]
            assert row["number_ot"] not in numbers
            numbers.add(row["number_ot"])

    def test_planned_trainings_have_codes(self, demo_book):
        planned = [r for r in demo_book.table("Инструктажи").rows if r["status"] == "Запланирован"]
        assert planned
        for row in planned:
            assert len(str(row["code"])) == 6
            assert not row["number_ot"]

    def test_both_starting_states_present(self, demo_book):
        """Демо должно показывать и просрочку, и будущие сроки."""
        today = date.today()
        planned_dates = [
            r["plan_date"] for r in demo_book.table("Инструктажи").rows
            if r["status"] == "Запланирован" and isinstance(r["plan_date"], date)
        ]
        assert any(d < today for d in planned_dates), "нет просроченных"
        assert any(d > today for d in planned_dates), "нет будущих"

    def test_counters_match_numbers(self, demo_book):
        ot = max(
            int(r["number_ot"][2:]) for r in demo_book.table("Инструктажи").rows
            if r["number_ot"]
        )
        assert int(demo_book.org_value("counter_ot")) == ot

    def test_demo_survives_roundtrip(self, demo_book):
        fresh = excel_store.open_book(demo_book.path, data_dir=demo_book.data_dir)
        assert len(fresh.table("Работники").rows) == len(demo_book.table("Работники").rows)
        assert len(fresh.table("Инструктажи").rows) == len(demo_book.table("Инструктажи").rows)

    def test_demo_book_is_usable(self, demo_book):
        """На демо-данных должны работать основные сценарии, а не только отображение."""
        from ot_registry import registration, services

        planned = [
            r for r in demo_book.table("Инструктажи").rows
            if r["status"] == "Запланирован" and r["plan_date"] <= date.today()
        ]
        assert planned
        row = planned[0]
        result = registration.register_training(
            demo_book,
            row["id"],
            {"fact_date": date.today(), "instructor_id": row["instructor_id"], "paper_confirmed": True},
        )
        assert result["training"]["status"] == "Проведён"
        assert result["training"]["number_ot"].startswith("ОТ")


class TestMakeExcelScript:
    def test_script_creates_file(self, tmp_path):
        import importlib.util

        script = Path(__file__).resolve().parents[1] / "scripts" / "make_excel.py"
        spec = importlib.util.spec_from_file_location("make_excel", script)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)

        out = tmp_path / "generated.xlsx"
        code = module.main(["--demo", "--out", str(out)])
        assert code == 0
        assert out.exists()

        wb = load_workbook(out)
        assert wb["Работники"].max_row > 1
        assert wb["Инструктажи"].protection.sheet

    def test_script_respects_no_protect(self, tmp_path):
        import importlib.util

        script = Path(__file__).resolve().parents[1] / "scripts" / "make_excel.py"
        spec = importlib.util.spec_from_file_location("make_excel2", script)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)

        out = tmp_path / "open.xlsx"
        assert module.main(["--no-protect", "--out", str(out)]) == 0
        wb = load_workbook(out)
        assert not wb["Инструктажи"].protection.sheet
